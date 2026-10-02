# -*- coding: utf-8 -*-
"""Utilidades para generar documentos Word con formato APA 7 (adaptado)."""

import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
BLACK = RGBColor(0x00, 0x00, 0x00)


def _style_run(run, size=12, bold=False, italic=False, color=BLACK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    return run


_TOKEN = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*]+?\*)")


def _add_markup(paragraph, text, size=12, bold=False, italic=False):
    """ admite **negrita**, *cursiva* y __negrita cursiva__ """
    for part in _TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _style_run(paragraph.add_run(part[2:-2]), size=size, bold=True, italic=italic)
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            _style_run(paragraph.add_run(part[2:-2]), size=size, bold=True, italic=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _style_run(paragraph.add_run(part[1:-1]), size=size, bold=bold, italic=True)
        else:
            _style_run(paragraph.add_run(part), size=size, bold=bold, italic=italic)


def _set_cell_borders(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        data = kwargs.get(edge, {"val": "nil"})
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), data.get("val", "single"))
        if data.get("val", "single") != "nil":
            el.set(qn("w:sz"), str(data.get("sz", 6)))
            el.set(qn("w:color"), data.get("color", "000000"))
        borders.append(el)
    tcPr.append(borders)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _add_page_field(paragraph):
    run = paragraph.add_run()
    _style_run(run, size=12)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


class Doc:
    def __init__(self, title, subtitle=None, author=None, affiliation=None,
                 date=None, extra=None):
        self.doc = Document()
        self._setup_styles()
        self._setup_page()
        self.table_index = 0
        self.figure_index = 0
        if title:
            self.title_page(title, subtitle, author, affiliation, date, extra)

    # ------------------------------------------------------------------ setup
    def _setup_styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.font.color.rgb = BLACK
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), FONT)

        specs = {
            "Heading 1": dict(size=12, bold=True, italic=False, align="center",
                              before=18, after=6),
            "Heading 2": dict(size=12, bold=True, italic=False, align="left",
                              before=12, after=6),
            "Heading 3": dict(size=12, bold=True, italic=True, align="left",
                              before=12, after=6),
            "Heading 4": dict(size=12, bold=True, italic=True, align="left",
                              before=12, after=6),
        }
        for name, spec in specs.items():
            s = self.doc.styles[name]
            s.font.name = FONT
            s.font.size = Pt(spec["size"])
            s.font.bold = spec["bold"]
            s.font.italic = spec["italic"]
            s.font.color.rgb = BLACK
            s.paragraph_format.space_before = Pt(spec["before"])
            s.paragraph_format.space_after = Pt(spec["after"])
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if spec["align"] == "center"
                else WD_ALIGN_PARAGRAPH.LEFT)
            rpr = s.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(attr), FONT)

    def _setup_page(self):
        for section in self.doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            header = section.header
            header.is_linked_to_previous = False
            hp = header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            _add_page_field(hp)

    # ------------------------------------------------------------ title page
    def title_page(self, title, subtitle=None, author=None, affiliation=None,
                   date=None, extra=None):
        for _ in range(4):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        _add_markup(p, title, size=12, bold=True)
        if subtitle:
            p2 = self.doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_markup(p2, subtitle, size=12, bold=True)
        self.doc.add_paragraph()
        for line in (author or []) + (affiliation or []) + (date or []):
            pa = self.doc.add_paragraph()
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_markup(pa, line, size=12)
        if extra:
            self.doc.add_paragraph()
            for line in extra:
                pe = self.doc.add_paragraph()
                pe.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_markup(pe, line, size=12)
        self.page_break()

    # ------------------------------------------------------------- headings
    def h1(self, text, new_page=True):
        if new_page and len(self.doc.paragraphs) > 1:
            self.page_break()
        p = self.doc.add_paragraph(style="Heading 1")
        p.add_run(text)
        return p

    def h2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        p.add_run(text)
        return p

    def h3(self, text):
        p = self.doc.add_paragraph(style="Heading 3")
        p.add_run(text)
        return p

    def h4(self, text):
        p = self.doc.add_paragraph(style="Heading 4")
        p.add_run(text)
        return p

    # ------------------------------------------------------------ paragraphs
    def p(self, text, size=12, align=None, spacing="double", indent=None,
          keep=False):
        par = self.doc.add_paragraph()
        if spacing == "single":
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            par.paragraph_format.space_after = Pt(6)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if indent is not None:
            par.paragraph_format.left_indent = Inches(indent)
        if keep:
            par.paragraph_format.keep_with_next = True
        _add_markup(par, text, size=size)
        return par

    def numbered(self, num, text, size=12):
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.5)
        par.paragraph_format.first_line_indent = Inches(-0.5)
        _add_markup(par, f"{num}. {text}", size=size)
        return par

    def bullet(self, text, size=12):
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.5)
        par.paragraph_format.first_line_indent = Inches(-0.25)
        _add_markup(par, f"\u2022  {text}", size=size)
        return par

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def spacer(self, lines=1):
        for _ in range(lines):
            self.doc.add_paragraph()

    # ---------------------------------------------------------------- tables
    def table(self, title, headers, rows, widths=None, size=10,
              align_center_cols=(), caption_prefix="Tabla"):
        """Tabla con encabezado APA (número en negrita + título en cursiva)."""
        self.table_index += 1
        cap = self.doc.add_paragraph()
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.paragraph_format.space_before = Pt(12)
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.keep_with_next = True
        _style_run(cap.add_run(f"{caption_prefix} {self.table_index}"), size=size, bold=True)
        _style_run(cap.add_run(f"  {title}"), size=size, italic=True)

        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge, val, sz in (("top", "single", 12), ("bottom", "single", 12),
                              ("left", "nil", 0), ("right", "nil", 0),
                              ("insideH", "single", 4), ("insideV", "nil", 0)):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val)
            if val != "nil":
                el.set(qn("w:sz"), str(sz))
                el.set(qn("w:color"), "000000")
            borders.append(el)
        tblPr.append(borders)

        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            _style_run(par.add_run(h), size=size, bold=True)
            _shade(cell, "E7E6E6")
            _set_cell_borders(cell, top={"val": "single", "sz": 12},
                              bottom={"val": "single", "sz": 8},
                              left={"val": "nil"}, right={"val": "nil"})

        for row_data in rows:
            row = t.add_row()
            for i, val in enumerate(row_data):
                cell = row.cells[i]
                cell.text = ""
                par = cell.paragraphs[0]
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(2)
                if i in align_center_cols:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_markup(par, str(val), size=size)
                _set_cell_borders(cell, top={"val": "nil"},
                                  bottom={"val": "single", "sz": 4},
                                  left={"val": "nil"}, right={"val": "nil"})

        if widths:
            total = sum(widths)
            usable = 6.5
            for i, w in enumerate(widths):
                inches = Inches(usable * w / total)
                for row in t.rows:
                    row.cells[i].width = inches
        after = self.doc.add_paragraph()
        after.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        after.paragraph_format.space_after = Pt(6)
        return t

    # -------------------------------------------------------------- figures
    def figure(self, path, title, width=5.6, caption_prefix="Figura"):
        self.figure_index += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(12)
        p.add_run().add_picture(path, width=Inches(width))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.paragraph_format.space_after = Pt(12)
        _style_run(cap.add_run(f"{caption_prefix} {self.figure_index}"), size=10, bold=True)
        _style_run(cap.add_run(f"  {title}"), size=10, italic=True)

    # ------------------------------------------------------------- abstract
    def abstract(self, text, heading="Resumen"):
        self.h1(heading, new_page=False)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_markup(p, text)
        self.page_break()

    # ----------------------------------------------------------- references
    def references(self, items, heading="Referencias"):
        self.h1(heading)
        for item in items:
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.5)
            par.paragraph_format.first_line_indent = Inches(-0.5)
            _add_markup(par, item)

    def save(self, path):
        core = self.doc.core_properties
        self.doc.save(path)
        return path
